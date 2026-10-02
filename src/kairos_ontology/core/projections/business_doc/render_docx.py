# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Cnext.eu
"""Word document of a business validation pack, from facts and a validated narrative (DD-254).

The layout is fixed so successive domains and versions read the same way: A4 portrait for
text, every diagram on its own landscape page, Yes / No / Comment boxes on every row a
reviewer confirms. Measurements follow the reference build the specification generalises.

Determinism: the same facts and narrative give the same ``word/document.xml``; the
core-properties timestamps come from the narrative date. python-docx is the optional extra
``business-doc``; a missing one is reported with the install command.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
from dataclasses import dataclass, field
from typing import Any, Callable, Optional

from .diagram_svg import CARDINALITY_MEANING, render_svg
from .facts import BusinessDocError
from .layout import LinkSpec, NodeSpec, layout_figure
from .narrative import Aliases, figure_relationships

NAVY = "1F3864"
GREY = "808080"
CHECK = "☐"

Rasterizer = Callable[[str, float, float], Optional[bytes]]


def _require_docx():
    try:
        import docx  # noqa: F401
    except ImportError as exc:
        raise BusinessDocError(
            "rendering a business validation document needs python-docx: "
            "uv sync --extra business-doc (a hub scaffolded before toolkit 5.26.1 has no "
            "such extra; use uv sync --extra documents, which also installs python-docx), "
            "or pip install 'kairos-ontology-toolkit[business-doc]'."
        ) from exc
    return docx


@dataclass
class RenderResult:
    docx: bytes
    svgs: dict[str, str] = field(default_factory=dict)  # file name -> SVG text
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------------------
# Figures.
# --------------------------------------------------------------------------------------
def _domain_title(name: str) -> str:
    return name.replace("-", " ").replace("_", " ").strip().capitalize()


def domain_view_figure(facts: dict[str, Any]):
    """Figure 1: the domain in the centre, each neighbour domain with what it is master of."""
    labels = [e["label"] for e in facts["entities"]]
    lines = tuple(labels[:8]) + ((f"... and {len(labels) - 8} more",) if len(labels) > 8 else ())
    centre = NodeSpec(
        "domain:" + facts["domain"], f"{_domain_title(facts['domain'])} domain", None, lines, "core"
    )
    nodes = [centre]
    links = []
    for row in facts["neighbour_domains"]:
        key = "domain:" + row["domain"]
        master = tuple(row["master_of"][:6])
        nodes.append(
            NodeSpec(key, f"{_domain_title(row['domain'])} domain", None, master, "external")
        )
        links.append(LinkSpec("N-" + row["domain"], key, centre.key, None, None, "links to"))
    return layout_figure(nodes, links, legend=False)


def model_figure(facts: dict[str, Any], figure: dict[str, Any], rel_ids: list[str]):
    """A logical-model figure: its entities, the relationships it holds, named gaps."""
    aliases = Aliases(facts)
    core = {e["iri"]: e for e in facts["entities"]}
    external = {e["iri"]: e for e in facts["externals"]}
    nodes: list[NodeSpec] = []
    for token in figure.get("entities") or ():
        key = aliases.node(token)
        if key in core:
            nodes.append(NodeSpec(key, core[key]["label"], None, (), "core"))
        elif key in external:
            ext = external[key]
            nodes.append(
                NodeSpec(
                    key, ext["label"], f"{_domain_title(ext['domain'])} domain", (), "external"
                )
            )
    links = []
    rels = {r["id"]: r for r in facts["relationships"]}
    for rel_id in rel_ids:
        rel = rels[rel_id]
        links.append(
            LinkSpec(
                rel_id,
                rel["from"],
                rel["to"],
                rel["from_card"],
                rel["to_card"],
                f"{rel_id} {rel['role_label']}",
            )
        )
    gaps = {str(g.get("id")): g for g in figure.get("_gaps", ())}
    for entry in figure.get("gaps") or ():
        gap_id = str(entry.get("id") if isinstance(entry, dict) else entry)
        gap = gaps.get(gap_id, {})
        nodes.append(
            NodeSpec(
                "gap:" + gap_id, str(gap.get("concept") or gap_id), None, (f"Gap {gap_id}",), "gap"
            )
        )
        near = aliases.node(entry.get("near")) if isinstance(entry, dict) else None
        if near:
            links.append(LinkSpec("G-" + gap_id, "gap:" + gap_id, near, None, None, "", True))
    return layout_figure(nodes, links)


# --------------------------------------------------------------------------------------
# Document helpers (python-docx).
# --------------------------------------------------------------------------------------
class _Writer:
    def __init__(self, docx_module, footer: str):
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm, Pt, RGBColor

        self.docx = docx_module
        self.Cm, self.Pt, self.RGB = Cm, Pt, RGBColor
        self.WD_ORIENT = WD_ORIENT
        self.doc = docx_module.Document()
        self.footer = footer
        self._styles()
        self._portrait(self.doc.sections[0])

    # ---- page setup ----------------------------------------------------------------
    def _styles(self) -> None:
        from docx.oxml.ns import qn

        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = self.Pt(10)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        for level, size in ((1, 16), (2, 12.5), (3, 11)):
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = "Calibri"
            style.font.size = self.Pt(size)
            style.font.bold = True
            style.font.color.rgb = self.RGB.from_string(NAVY)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = self.Pt(12 if level == 1 else 8)
            style.paragraph_format.space_after = self.Pt(4)

    def _footer(self, section) -> None:
        paragraph = section.footer.paragraphs[0]
        paragraph.text = ""
        run = paragraph.add_run(self.footer)
        run.font.size = self.Pt(8)
        run.font.color.rgb = self.RGB.from_string(GREY)

    def _portrait(self, section) -> None:
        section.orientation = self.WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = self.Cm(21.0), self.Cm(29.7)
        section.left_margin = section.right_margin = self.Cm(2.0)
        section.top_margin = section.bottom_margin = self.Cm(1.8)
        self._footer(section)

    def _landscape(self, section) -> None:
        section.orientation = self.WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = self.Cm(29.7), self.Cm(21.0)
        section.left_margin = section.right_margin = self.Cm(2.0)
        section.top_margin = section.bottom_margin = self.Cm(1.4)
        self._footer(section)

    def new_section(self, landscape: bool) -> None:
        from docx.enum.section import WD_SECTION

        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        (self._landscape if landscape else self._portrait)(section)

    # ---- text ----------------------------------------------------------------------
    def rich(
        self,
        paragraph,
        text: str,
        *,
        size: Optional[float] = None,
        color=None,
        bold: bool = False,
        italic: bool = False,
    ) -> None:
        """Append *text* to *paragraph*; ``**bold**`` spans are honoured."""
        for part in re.split(r"(\*\*[^*]+\*\*)", str(text)):
            if not part:
                continue
            strong = part.startswith("**") and part.endswith("**") and len(part) > 4
            run = paragraph.add_run(part[2:-2] if strong else part)
            run.bold = bold or strong
            run.italic = italic
            if size:
                run.font.size = self.Pt(size)
            if color:
                run.font.color.rgb = self.RGB.from_string(color)

    def heading(self, text: str, level: int) -> None:
        self.doc.add_heading(str(text), level=level)

    def para(self, text: str = "", **style) -> Any:
        paragraph = self.doc.add_paragraph()
        if text:
            self.rich(paragraph, text, **style)
        return paragraph

    def bullets(self, items: list[str]) -> None:
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Bullet")
            self.rich(paragraph, item)

    # ---- tables --------------------------------------------------------------------
    def _shade(self, cell, fill: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shading)

    def _left_border(self, cell, color: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            if side == "left":
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), "24")
                edge.set(qn("w:color"), color)
            else:
                edge.set(qn("w:val"), "nil")
            borders.append(edge)
        cell._tc.get_or_add_tcPr().append(borders)

    def _keep_together(self, table) -> None:
        """No row splits across pages, and each row stays on the page of the next one."""
        from docx.oxml import OxmlElement

        for index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if index < len(table.rows) - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True

    def _widths(self, table, widths: list[float]) -> None:
        table.autofit = False
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = self.Cm(width)

    def _cell(self, cell, text: str, *, size: float, bold=False, color=None) -> None:
        cell.text = ""
        self.rich(cell.paragraphs[0], text, size=size, bold=bold, color=color)

    def table(
        self,
        headers: Optional[list[str]],
        rows: list[list[str]],
        widths: list[float],
        *,
        font: float = 9,
        header_fill: str = NAVY,
        zebra: bool = True,
    ) -> Any:
        """A grid; ``headers=None`` for a plain key/value block with no header row."""
        offset = 1 if headers else 0
        table = self.doc.add_table(rows=offset + len(rows), cols=len(widths))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, headers or ()):
            self._cell(cell, text, size=font, bold=True, color="FFFFFF")
            self._shade(cell, header_fill)
        for index, row in enumerate(rows, start=offset):
            for cell, text in zip(table.rows[index].cells, row):
                self._cell(cell, text, size=font)
                if zebra and index % 2 == 0:
                    self._shade(cell, "F2F2F2")
        self._widths(table, widths)
        self.para()
        return table

    def fields(self, rows: list[list[str]]) -> None:
        self.table(
            ["Field", "Meaning", "Data type & format", "Code list", "Yes", "No", "Comment"],
            rows,
            [3.0, 5.2, 2.4, 2.4, 1.0, 1.0, 2.0],
        )

    def relations(self, rows: list[list[str]]) -> None:
        self.table(
            ["#", "Read in one direction", "Read in the other direction", "Yes", "No", "Comment"],
            rows,
            [0.9, 5.9, 5.8, 1.0, 1.0, 2.4],
        )

    def panel(
        self, title: str, items: list[str], *, fill: str = "FFF4D6", edge: str = "E0A030"
    ) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        self._shade(cell, fill)
        self._left_border(cell, edge)
        self.rich(cell.paragraphs[0], title, bold=True)
        for item in items:
            self.rich(cell.add_paragraph(), "• " + str(item))
        self._widths(table, [17.0])
        self.para()

    def definition(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        self._shade(cell, "EEF3F9")
        self._left_border(cell, NAVY)
        self.rich(cell.paragraphs[0], text)
        self._widths(table, [17.0])
        self.check("Is this definition correct and understandable?")

    def check(self, question: str) -> None:
        table = self.doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(
            table.rows[0].cells, [question, f"Yes {CHECK}", f"No {CHECK}", "Comment:"]
        ):
            self._cell(cell, text, size=9, bold=(text == question))
        self._widths(table, [7.0, 1.6, 1.6, 6.8])
        self.para()

    def missing_box(self, title: str) -> None:
        from docx.enum.table import WD_ROW_HEIGHT_RULE

        table = self.doc.add_table(rows=2, cols=1)
        table.style = "Table Grid"
        self._cell(table.rows[0].cells[0], title, size=9, bold=True)
        self._shade(table.rows[0].cells[0], "EEF3F9")
        table.rows[1].height = self.Cm(2.2)
        table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        self._keep_together(table)
        self._widths(table, [17.0])
        self.para()

    def figure(
        self,
        caption: str,
        intro: str,
        png: Optional[bytes],
        svg_name: str,
        width: float,
        height: float,
    ) -> Optional[str]:
        """A landscape page with *caption*, *intro* and the diagram; a warning when no PNG."""
        self.new_section(landscape=True)
        self.heading(caption, 2)
        if intro:
            self.para(intro)
        if png is None:
            self.para(
                f"[Diagram not rendered: open {svg_name} beside this document.]",
                italic=True,
                color=GREY,
            )
            return f"{caption}: no headless browser rendered the diagram; see {svg_name}"
        ratio = width / height if height else 1.0
        width_cm = min(25.7, 14.5 * ratio)
        self.doc.add_picture(io.BytesIO(png), width=self.Cm(width_cm))
        return None

    def save(self, title: str, date: _dt.date) -> bytes:
        stamp = _dt.datetime(date.year, date.month, date.day, tzinfo=_dt.timezone.utc)
        props = self.doc.core_properties
        props.author = props.last_modified_by = "kairos-ontology business-doc"
        props.title = title
        props.created = props.modified = props.last_printed = stamp
        props.revision = 1
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


# --------------------------------------------------------------------------------------
# The document.
# --------------------------------------------------------------------------------------
def _date(value: Any) -> _dt.date:
    if isinstance(value, _dt.date):
        return value
    return _dt.date.fromisoformat(str(value))


_DEFAULT_HOW_TO_USE = [
    "Tick **Yes** or **No** on every row. Use **Comment** for anything that is wrong, unclear "
    "or missing.",
    "Write comments anywhere in the document; every one is read.",
    "Concepts owned by another domain are only referenced here, and name the domain that owns "
    "them.",
]


def _cardinality_phrase(code: str) -> str:
    return CARDINALITY_MEANING[code]


def render_document(
    facts: dict[str, Any],
    narrative: dict[str, Any],
    *,
    version: Optional[int] = None,
    rasterize: Optional[Rasterizer] = None,
) -> RenderResult:
    """Render the document. *narrative* must already pass ``validate_narrative``."""
    docx = _require_docx()
    document = narrative.get("document") or {}
    number = int(version or document.get("version") or 1)
    date = _date(document.get("date"))
    domain_title = _domain_title(facts["domain"])
    company = str(document.get("company") or facts.get("hub") or "").strip()
    footer = (f"{company} – " if company else "") + (
        f"{domain_title} domain – logical model for business validation (v{number})"
    )
    w = _Writer(docx, footer)
    result = RenderResult(docx=b"")
    aliases = Aliases(facts)
    figure_number = 0

    def draw(caption: str, intro: str, layout, stem: str) -> None:
        nonlocal figure_number
        figure_number += 1
        svg = render_svg(layout)
        name = f"{facts['domain']}-validation-v{number}-{stem}.svg"
        result.svgs[name] = svg
        png = rasterize(svg, layout.width, layout.height) if rasterize else None
        warning = w.figure(caption, intro, png, name, layout.width, layout.height)
        if warning:
            result.warnings.append(warning)

    # ---- title block ------------------------------------------------------------------
    w.para(f"{domain_title} domain", size=24, bold=True, color=NAVY)
    w.para("Logical data model for business validation", size=14, color=NAVY)
    meta = [
        ["Document version", str(number)],
        ["Date", date.isoformat()],
        ["Model version", facts.get("model_version") or "not set"],
        ["Facts", facts["facts_hash"][:12]],
    ]
    if document.get("previous_version"):
        meta.append(["Previous version", str(document["previous_version"])])
    w.table(None, meta, [5.0, 12.0], zebra=False)
    if document.get("ai_drafted"):
        w.para("AI-drafted, not user-confirmed.", bold=True, color="C0504D")

    w.heading("How to use this document", 2)
    w.bullets([str(i) for i in (narrative.get("how_to_use") or _DEFAULT_HOW_TO_USE)])
    if narrative.get("what_changed"):
        previous = document.get("previous_version")
        title = f"What changed since version {previous}" if previous else "What changed"
        w.panel(title, [str(i) for i in narrative["what_changed"]], fill="EEF3F9", edge=NAVY)

    # ---- 1 terms ----------------------------------------------------------------------
    w.heading("1 Terms used in this document", 1)
    glossary = sorted(narrative.get("glossary") or (), key=lambda g: str(g["term"]).lower())
    if glossary:
        w.table(
            ["Term", "Meaning in this document", "Owned by"],
            [[str(g["term"]), str(g["meaning"]), str(g.get("owned_by") or "")] for g in glossary],
            [4.0, 9.5, 3.5],
        )
    else:
        w.para("No terms are defined yet.", italic=True)

    # ---- 2 domain view ----------------------------------------------------------------
    w.heading(f"2 The {facts['domain']} domain and its neighbours", 1)
    view = narrative.get("domain_view") or {}
    w.para(
        f"Each neighbour domain is **master of its data**: it links to the "
        f"{facts['domain']} domain, it does not copy into it. Figure 1 is on the next page."
    )
    rows = []
    for row in facts["neighbour_domains"]:
        how = (view.get("rows") or {}).get(row["domain"]) or (
            "Linked through " + ", ".join(row["relationships"])
        )
        rows.append(
            [_domain_title(row["domain"]), ", ".join(row["master_of"]), str(how), CHECK, CHECK, ""]
        )
    if rows:
        w.table(
            [
                "Domain",
                "Master of",
                f"How it is related to the {facts['domain']}",
                "Yes",
                "No",
                "Comment",
            ],
            rows,
            [3.0, 4.0, 5.4, 1.0, 1.0, 2.6],
        )
    else:
        w.para("No other domain is linked to this one yet.", italic=True)
    if view.get("why"):
        w.panel("Why it works this way", [str(i) for i in view["why"]])
    w.check("Do these domains and relationships match how the business works?")
    draw(
        f"Figure 1: The {facts['domain']} domain and its neighbours",
        "",
        domain_view_figure(facts),
        "figure-1",
    )

    # ---- 3 logical data model ---------------------------------------------------------
    w.new_section(landscape=False)
    w.heading("3 Logical data model", 1)
    w.para(
        "Each line between two boxes is a relationship. The marks at each end of a line say "
        "how many objects at that end belong to one object at the other end."
    )
    legend_rows = []
    for code, shape in (
        ("1", "Two bars"),
        ("01", "Circle and bar"),
        ("1n", "Bar and crow's foot"),
        ("0n", "Circle and crow's foot"),
    ):
        example = next((r for r in facts["relationships"] if r["to_card"] == code), None)
        if example is not None:
            frm = aliases.node(example["from"])
            to = aliases.node(example["to"])
            text = (
                f"Each {_label_of(facts, frm)} has {_cardinality_phrase(code)} "
                f"{_label_of(facts, to)}."
            )
        else:
            text = ""
        legend_rows.append([shape, _cardinality_phrase(code).capitalize(), text])
    w.table(["Line end", "Meaning", "Example"], legend_rows, [4.0, 3.5, 9.5])

    narrative_figures = list(narrative.get("figures") or ())
    placed = figure_relationships(facts, narrative)
    gaps = list(narrative.get("gaps") or ())
    for figure in narrative_figures:
        figure = dict(figure, _gaps=gaps)
        rel_ids = placed.get(str(figure.get("id")), [])
        draw(
            f"Figure {figure['id']}: {figure['title']}",
            str(figure.get("intro") or ""),
            model_figure(facts, figure, rel_ids),
            f"figure-{figure['id']}",
        )
    w.new_section(landscape=False)
    wording = narrative.get("relationships") or {}
    rels = {r["id"]: r for r in facts["relationships"]}
    for figure in narrative_figures:
        rel_ids = placed.get(str(figure.get("id")), [])
        if not rel_ids:
            continue
        w.heading(f"Relationships in Figure {figure['id']}", 3)
        w.relations(
            [
                [rid, str(wording[rid]["one"]), str(wording[rid]["other"]), CHECK, CHECK, ""]
                for rid in rel_ids
                if rid in rels and rid in wording
            ]
        )
    w.missing_box(
        "Anything missing from the logical data model? (entities, relationships, cardinalities)"
    )

    # ---- 4..n entities ----------------------------------------------------------------
    blocks = {aliases.entity(k): v for k, v in (narrative.get("entities") or {}).items()}
    section = 3
    for entity in facts["entities"]:
        block = blocks.get(entity["iri"])
        if block is None:
            continue  # omitted with a reason (validated)
        section += 1
        _entity_section(w, facts, entity, block, section, aliases)

    # ---- gaps and decisions -----------------------------------------------------------
    section += 1
    w.heading(f"{section} Gaps and open decisions", 1)
    w.heading("Concepts the business uses that the model does not hold yet", 3)
    if gaps:
        w.table(
            [
                "#",
                "Concept",
                "What the model holds today",
                f"Needed? (Yes {CHECK} No {CHECK})",
                "Comment",
            ],
            [
                [
                    str(g["id"]),
                    str(g.get("concept") or ""),
                    str(g.get("today") or ""),
                    f"Yes {CHECK}  No {CHECK}",
                    "",
                ]
                for g in gaps
            ],
            [1.0, 4.0, 6.0, 2.6, 3.4],
        )
    else:
        w.para("No gaps are listed.", italic=True)
    w.heading("Decisions needed from the business", 3)
    decisions = list(narrative.get("decisions") or ())
    if decisions:
        w.table(
            ["#", "Question", "Current behaviour", "Your answer"],
            [
                [str(d["id"]), str(d.get("question") or ""), str(d.get("current") or ""), ""]
                for d in decisions
            ],
            [1.0, 6.0, 5.0, 5.0],
        )
    else:
        w.para("No decisions are open.", italic=True)

    # ---- confirmation -----------------------------------------------------------------
    section += 1
    w.heading(f"{section} Confirmation", 1)
    w.para(
        "By signing, you confirm that this document, with your answers and comments, "
        "describes how the business works."
    )
    w.table(
        ["Name", "Role", "Date", "Signature"],
        [["", "Business representative", "", ""], ["", "Data architect", "", ""]],
        [5.0, 4.0, 3.0, 5.0],
        zebra=False,
    )
    w.missing_box("General comments")

    remarks = list(narrative.get("review_remarks") or ())
    if remarks:
        previous = document.get("previous_version") or (number - 1)
        w.heading("Appendix A Review remarks and answers", 1)
        w.table(
            ["#", f"Remark on version {previous}", f"How version {number} answers it"],
            [
                [str(r.get("id")), str(r.get("remark") or ""), str(r.get("answer") or "")]
                for r in remarks
            ],
            [1.0, 8.0, 8.0],
        )

    result.docx = w.save(
        f"{domain_title} domain - logical data model for business validation (v{number})", date
    )
    return result


def _label_of(facts: dict[str, Any], iri: Optional[str]) -> str:
    for collection in ("entities", "externals"):
        for item in facts[collection]:
            if item["iri"] == iri:
                return item["label"].lower()
    return str(iri)


def _entity_section(w: _Writer, facts, entity, block, number: int, aliases: Aliases) -> None:
    label = entity["label"]
    w.heading(f"{number} {block.get('heading') or label}", 1)
    w.definition(str(block["definition"]))
    if block.get("why"):
        w.panel("Why it works this way", [str(i) for i in block["why"]])

    w.heading("Identification", 3)
    ident = entity["identification"]
    rows = [
        [
            f"{ident['technical_id']['column']} (PK)",
            "Technical identifier the platform generates; not shown to users.",
            ident["technical_id"]["type"],
            CHECK,
            CHECK,
            "",
        ]
    ]
    key_columns = ident["business_key"]["columns"]
    if key_columns:
        sources = ", ".join(entity["sources"]) or "its source systems"
        rows.append(
            [
                ", ".join(key_columns) + " (BK)",
                str(
                    block.get("business_key") or f"Identifies one {label.lower()} within {sources}."
                ),
                "Text",
                CHECK,
                CHECK,
                "",
            ]
        )
    w.table(
        ["Field", "Meaning", "Data type & format", "Yes", "No", "Comment"],
        rows,
        [3.6, 6.4, 2.8, 1.0, 1.0, 2.2],
    )

    fields = {f["property"]: f for f in entity["fields"]}
    meanings = {
        aliases.field(entity["iri"], k): v for k, v in (block.get("field_meanings") or {}).items()
    }
    held = {
        aliases.field(entity["iri"], k): v for k, v in (block.get("held_references") or {}).items()
    }
    if held:
        w.heading(f"Numbers of other objects, still held on the {label.lower()}", 3)
        w.table(
            ["Field", "What it is the number of", "Yes", "No", "Comment"],
            [
                [fields[p]["label"], str(text), CHECK, CHECK, ""]
                for p, text in held.items()
                if p in fields
            ],
            [3.6, 9.2, 1.0, 1.0, 2.2],
        )
    for group in block.get("field_groups") or ():
        props = [aliases.field(entity["iri"], t) for t in group.get("properties") or ()]
        rows = []
        for prop in props:
            item = fields.get(prop)
            if item is None:
                continue
            meaning = meanings.get(prop) or item["comment"] or ""
            datatype = item["datatype"] + (", required" if item["required"] else "")
            code = item["code_list"]["label"] if item["code_list"] else ""
            rows.append([item["label"], str(meaning), datatype, code, CHECK, CHECK, ""])
        if rows:
            w.heading(str(group["title"]), 3)
            w.fields(rows)
    w.para(
        f"Links from the {label.lower()} to other objects are relationships (section 3), "
        "not fields.",
        italic=True,
        color=GREY,
    )
    w.missing_box(f"Anything missing from the {label.lower()}?")


# --------------------------------------------------------------------------------------
# PDF preview (best effort: LibreOffice, else Word on Windows).
# --------------------------------------------------------------------------------------
def pdf_preview(docx_path) -> Optional[str]:
    """Write ``<docx>.pdf`` beside *docx_path*; return its path, or ``None`` when no converter.

    The PDF is a preview for the reviewer's convenience. It is never compared or gated, so a
    converter's own variation does not matter.
    """
    import os
    import shutil
    import subprocess
    from pathlib import Path

    source = Path(docx_path).resolve()
    target = source.with_suffix(".pdf")
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None and os.name == "nt":
        for base in (os.environ.get("PROGRAMFILES"), os.environ.get("PROGRAMFILES(X86)")):
            candidate = Path(base or "") / "LibreOffice" / "program" / "soffice.exe"
            if base and candidate.is_file():
                office = str(candidate)
                break
    try:
        if office:
            subprocess.run(
                [
                    office,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(source.parent),
                    str(source),
                ],
                check=True,
                capture_output=True,
                timeout=180,
            )
        elif os.name == "nt" and shutil.which("powershell"):
            script = (
                "$ErrorActionPreference='Stop';"
                "$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
                "try{$d=$w.Documents.Open($env:KAIROS_DOCX,$false,$true);"
                "$d.ExportAsFixedFormat($env:KAIROS_PDF,17);$d.Close($false)}"
                "finally{$w.Quit()}"
            )
            env = dict(os.environ, KAIROS_DOCX=str(source), KAIROS_PDF=str(target))
            subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                check=True,
                capture_output=True,
                timeout=180,
                env=env,
            )
        else:
            return None
    except (OSError, subprocess.SubprocessError):
        return None
    return str(target) if target.is_file() else None
