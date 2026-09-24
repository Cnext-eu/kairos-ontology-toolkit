# SPDX-License-Identifier: Apache-2.0
# Copyright 2026 Cnext.eu
"""Text from staged business-discovery documents (#907).

DD-233 fails `validate` until every staged document has an extraction; on one hub 20 of
31 were Office formats nothing in the toolkit could open. These pin the readers against
real files built here, not mocks, and the one folder that opts a file out: `.import/drafts/`.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from click.testing import CliRunner

from kairos_ontology.cli.main import cli
from kairos_ontology.core.document_text import (
    BUSINESS_DOCUMENT_SUFFIXES,
    DocumentReadError,
    read_document_text,
)
from kairos_ontology.core.import_evidence import audit_import_evidence, find_misplaced_documents


def _docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Operating model", level=1)
    document.add_paragraph("A consignment is booked by the forwarder.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Term", "Meaning"
    table.cell(1, 0).text, table.cell(1, 1).text = "Laytime", "Time allowed to load"
    document.save(str(path))
    return path


def _pptx(path: Path) -> Path:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Our services"
    slide.placeholders[1].text = "Door-to-door freight"
    slide.notes_slide.notes_text_frame.text = "Mention the reefer fleet."
    deck.slides.add_slide(deck.slide_layouts[6])  # blank: no text at all
    deck.save(str(path))
    return path


def _xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Lane", "Rate"])
    sheet.append(["ANR-RTM", 120])
    workbook.save(str(path))
    return path


def _pdf(path: Path) -> Path:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)
    return path


class TestReaders:
    def test_docx_keeps_headings_paragraphs_and_tables(self, tmp_path):
        doc = read_document_text(_docx(tmp_path / "model.docx"))
        assert doc.format == "docx"
        assert "## Operating model" in doc.text
        assert "A consignment is booked by the forwarder." in doc.text
        assert "| Laytime | Time allowed to load |" in doc.text

    def test_pptx_reads_slides_and_notes_and_names_image_only_slides(self, tmp_path):
        doc = read_document_text(_pptx(tmp_path / "deck.pptx"))
        assert doc.units == 2
        assert "## Slide 1" in doc.text and "Door-to-door freight" in doc.text
        assert "Speaker notes: Mention the reefer fleet." in doc.text
        assert any("1 slide(s) carry no text" in note for note in doc.notes)

    def test_xlsx_reads_every_sheet_as_rows(self, tmp_path):
        doc = read_document_text(_xlsx(tmp_path / "rates.xlsx"))
        assert "## Sheet: Rates" in doc.text
        assert "| ANR-RTM | 120 |" in doc.text

    def test_a_pdf_without_a_text_layer_says_so(self, tmp_path):
        doc = read_document_text(_pdf(tmp_path / "scan.pdf"))
        assert doc.units == 1
        assert any("no text layer" in note for note in doc.notes)

    def test_html_is_reduced_to_its_text(self, tmp_path):
        path = tmp_path / "about.htm"
        path.write_text(
            "<html><style>p{}</style><body><h1>About</h1><p>We ship &amp; store.</p></body>",
            encoding="utf-8",
        )
        text = read_document_text(path).text
        assert "About" in text and "We ship & store." in text
        assert "<" not in text and "p{}" not in text

    @pytest.mark.parametrize("suffix, modern", [(".doc", ".docx"), (".ppt", ".pptx")])
    def test_a_legacy_format_says_what_to_convert_it_to(self, tmp_path, suffix, modern):
        path = tmp_path / f"old{suffix}"
        path.write_bytes(b"\xd0\xcf\x11\xe0")
        with pytest.raises(DocumentReadError, match=modern):
            read_document_text(path)

    def test_an_image_points_at_visual_reading_or_drafts(self, tmp_path):
        path = tmp_path / "flow.png"
        path.write_bytes(b"\x89PNG")
        with pytest.raises(DocumentReadError, match="drafts"):
            read_document_text(path)

    def test_a_missing_extra_names_the_install_command(self, tmp_path, monkeypatch):
        import builtins

        real_import = builtins.__import__

        def no_docx(name, *args, **kwargs):
            if name == "docx":
                raise ImportError("no docx")
            return real_import(name, *args, **kwargs)

        path = _docx(tmp_path / "model.docx")
        monkeypatch.setattr(builtins, "__import__", no_docx)
        with pytest.raises(DocumentReadError, match="uv sync --extra documents"):
            read_document_text(path)


class TestTheCommand:
    def test_read_document_prints_the_text(self, tmp_path):
        result = CliRunner().invoke(cli, ["read-document", str(_docx(tmp_path / "m.docx"))])
        assert result.exit_code == 0, result.output
        assert "A consignment is booked by the forwarder." in result.output

    def test_json_carries_format_units_and_notes(self, tmp_path):
        import json

        result = CliRunner().invoke(
            cli, ["read-document", str(_pptx(tmp_path / "d.pptx")), "--format", "json"]
        )
        payload = json.loads(result.stdout)
        assert payload["format"] == "pptx" and payload["units"] == 2 and payload["notes"]

    def test_an_unreadable_file_is_a_clean_error(self, tmp_path):
        path = tmp_path / "old.xls"
        path.write_bytes(b"x")
        result = CliRunner().invoke(cli, ["read-document", str(path)])
        assert result.exit_code != 0
        assert ".xlsx" in result.output


class TestDraftsAndTheSharedDefinition:
    def _hub(self, tmp_path):
        (tmp_path / ".import").mkdir()
        hub = tmp_path / "ontology-hub"
        (hub / "businessdiscovery" / "_extractions").mkdir(parents=True)
        (hub / "integration" / "discovery" / "bi").mkdir(parents=True)
        return hub

    def test_a_file_in_drafts_is_neither_misplaced_nor_owed_an_extraction(self, tmp_path):
        hub = self._hub(tmp_path)
        drafts = tmp_path / ".import" / "drafts"
        drafts.mkdir()
        (drafts / "old-deck.pptx").write_bytes(b"x")

        assert audit_import_evidence(tmp_path, hub).findings == []

    def test_every_file_in_businessdiscovery_is_owed_an_extraction(self, tmp_path):
        """Whatever its suffix: the folder is the statement of intent."""
        hub = self._hub(tmp_path)
        bd = tmp_path / ".import" / "businessdiscovery"
        bd.mkdir()
        for name in ("notes.txt", "flow.png", "partners.xml"):
            (bd / name).write_bytes(b"x")

        kinds = [f.kind for f in audit_import_evidence(tmp_path, hub).findings]
        assert kinds == ["unextracted"] * 3

    def test_an_htm_staged_elsewhere_is_reported_as_misplaced(self, tmp_path):
        """#907's 27-vs-31: .xml/.htm were processed by discovery but invisible here."""
        stray = tmp_path / "Input"
        stray.mkdir()
        (stray / "about.htm").write_text("x", encoding="utf-8")
        (stray / "partners.xml").write_text("x", encoding="utf-8")

        assert [p.name for p in find_misplaced_documents(tmp_path)] == [
            "about.htm",
            "partners.xml",
        ]

    def test_every_readable_business_format_counts_as_a_business_document(self):
        from kairos_ontology.core.document_text import READERS

        assert {s for s in READERS if s not in {".csv", ".txt"}} <= BUSINESS_DOCUMENT_SUFFIXES
